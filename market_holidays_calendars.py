import argparse

from datetime import date
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment

import pandas as pd
import pandas_market_calendars as mcal

# =====================================================================
# 1. Discover supported exchange calendars
# =====================================================================

def list_supported_exchanges() -> tuple[list[str], list[dict[str, str]]]:
    """
    Find the exchange calendars available in the installed library.

    Return one name per calendar and record any calendars
    that could not be loaded.
    """

    supported_exchanges = []
    failed_exchanges = []
    seen_calendars = set()

    for exchange_name in mcal.get_calendar_names():
        try:
            calendar = mcal.get_calendar(exchange_name)
        except Exception as exc:
            failed_exchanges.append(
                {
                    "Exchange": exchange_name,
                    "Error": f"{type(exc).__name__}: {exc}"
                }
            )
            continue

        # The library may register multiple names for the same calendar.
        canonical_name = calendar.name
        if canonical_name in seen_calendars:
            continue

        seen_calendars.add(canonical_name)

        # Keeping the registered name so the library can load it again.
        supported_exchanges.append(exchange_name)

    return supported_exchanges, failed_exchanges


# ===========================================================
# 2. Find exchange holidays and early closes
# ===========================================================
 
def get_market_holidays(
    exchange: str,
    start_date: str,
    end_date: str,
    include_early_closes: bool = True
) -> dict:
    """
    Find full-day closures and optional early closes for an exchange.
    The start and end dates are both included in the search.

    Parameters
    ----------
    exchange: A calendar name supported by pandas_market_calendars,
    start_date: Inclusive start date in YYYY-MM-DD format.
    end_date: Inclusive end date in YYYY-MM-DD format.
    include_early_closes: Whether to include sessions with an early market close.
    """

    if not isinstance(exchange, str) or not exchange.strip():
        raise ValueError("Exchange must be a non-empty string.")


    # Allowing exchange names to be entered without worrying about capitalization.
    exchange_lookup = {
        name.casefold(): name for name in mcal.get_calendar_names()
    }

    exchange_key = exchange.strip().casefold()

    if exchange_key not in exchange_lookup:
        raise ValueError(
            f"Unsupported exchange: {exchange}. "
            "Use a calendar name supported by pandas_market_calendars."
        )

    calendar_name = exchange_lookup[exchange_key]
    calendar = mcal.get_calendar(calendar_name)

    # Checking the date format and making sure the range runs forward.
    try:
        start = date.fromisoformat(start_date)
        end = date.fromisoformat(end_date)
    except (TypeError, ValueError) as exc:
        raise ValueError("Dates must use YYYY-MM-DD format.") from exc

    if (start.isoformat() != start_date or end.isoformat() != end_date):
        raise ValueError("Dates must use YYYY-MM-DD format.")

    if start > end:
        raise ValueError("Start date cannot be after end date.")

    # Generating dates on which this exchange would normally be scheduled to trade.
    expected_weekdays = pd.bdate_range(
        start=start,
        end=end,
        freq="C",
        weekmask=calendar.weekmask
    )

    # Retrieving actual exchange trading sessions.
    valid_sessions = pd.DatetimeIndex(
        calendar.valid_days(start_date=start, end_date=end)
    )

    # Removing timezone information so both date indexes can be compared consistently.
    valid_sessions = valid_sessions.tz_localize(None).normalize()

    # Identifying full-day market closures.
    holiday_dates = expected_weekdays.difference(valid_sessions)
    holidays = [day.date().isoformat() for day in holiday_dates]

    # Identifying early-close trading sessions.
    early_closes = []

    if include_early_closes:
        schedule = calendar.schedule(start_date=start, end_date=end)
        early_close_schedule = calendar.early_closes(schedule)
        early_closes = [
            day.date().isoformat() for day in early_close_schedule.index
        ]

    # Official Cboe 2026 calendar lists Dec 24 as a Christmas early close and this one documented exception only when it falls in the date range
    if include_early_closes and (
        calendar_name == "CBOE_Index_Options"
        or calendar.name == "CBOE_Index_Options"
    ):
        christmas_eve = date(2026, 12, 24)
        if start <= christmas_eve <= end:
            christmas_eve_text = christmas_eve.isoformat()
            if christmas_eve_text not in holidays:
                early_closes = sorted(set(early_closes) | {christmas_eve_text})

    # Returning dates as strings so they are easy to use in other outputs.
    return {
        "exchange": calendar_name,
        "start_date": start.isoformat(),
        "end_date": end.isoformat(),
        "holidays": holidays,
        "early_closes": early_closes
    }

# =====================================================================
# 3. Run the calendar from the command line
# =====================================================================

def main() -> None:
    """
    Read command-line options and save the requested calendars to Excel.
    """

    # Let the user choose exchanges, dates, and the output workbook
    parser = argparse.ArgumentParser(
        description="Generate an Excel market holiday calendar for one or more exchanges."
    )

    # Choosing specific exchanges or all supported exchanges.
    exchange_group = parser.add_mutually_exclusive_group(required=True)
    exchange_group.add_argument(
        "--exchange",
        nargs="+",
        help="Exchange calendar names, e.g. XNYS CBOE_Index_Options.",
    )
    exchange_group.add_argument(
        "--all",
        action="store_true",
        help="Generate calendars for all supported exchanges.",
    )

    parser.add_argument(
        "--start", required=True, help="Inclusive start date in YYYY-MM-DD format."
    )
    parser.add_argument(
        "--end", required=True, help="Inclusive end date in YYYY-MM-DD format."
    )
    parser.add_argument(
        "--no-early-closes",
        action="store_true",
        help="Exclude early-close sessions.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=(
            Path(__file__).resolve().parent
            / "output"
            / "market_holidays_output.xlsx"
        ),
        help="Destination Excel workbook.",
    )

    args = parser.parse_args()
    output_path = args.output

    if output_path.suffix.lower() != ".xlsx":
        parser.error("Output filename must end with .xlsx")

    if args.all:
        # Check dates once so a bad range is not logged as failures.
        try:
            start = date.fromisoformat(args.start)
            end = date.fromisoformat(args.end)
        except (TypeError, ValueError):
            parser.error("Dates must use YYYY-MM-DD format.")

        if start.isoformat() != args.start or end.isoformat() != args.end:
            parser.error("Dates must use YYYY-MM-DD format.")
        if start > end:
            parser.error("Start date cannot be after end date.")

        exchange_names, failed_exchanges = list_supported_exchanges()
    else:
        exchange_names = args.exchange
        failed_exchanges = []


    # Collecting one summary row per exchange and one row per calendar event.
    summary_records = []
    event_records = []

    # Generate the holiday calendar for each selected exchange.
    for exchange_name in exchange_names:
        try:
            result = get_market_holidays(
                exchange=exchange_name,
                start_date=args.start,
                end_date=args.end,
                include_early_closes=not args.no_early_closes
            )
        except Exception as exc:
            # When a specific exchange is requested, report the error.
            if not args.all:
                if isinstance(exc, ValueError):
                    parser.error(str(exc))

                raise

            # With --all, record the failure and keep going.
            failed_exchanges.append(
                {
                    "Exchange": exchange_name,
                    "Error": f"{type(exc).__name__}: {exc}"
                }
            )
            continue

        # Add one summary row for each successful exchange.
        summary_records.append(
            {
                "Exchange": result["exchange"],
                "Start Date": date.fromisoformat(result["start_date"]),
                "End Date": date.fromisoformat(result["end_date"]),
                "Full-day Holidays": len(result["holidays"]),
                "Early Closes": len(result["early_closes"]),
            }
        )

        # Add the full-day closures.
        for holiday in result["holidays"]:
            event_records.append(
                {
                    "Exchange": result["exchange"],
                    "Date": date.fromisoformat(holiday),
                    "Event Type": "Full-day Holiday",
                }
            )

        # Adding the early-close trading sessions.
        for early_close in result["early_closes"]:
            event_records.append(
                {
                    "Exchange": result["exchange"],
                    "Date": date.fromisoformat(early_close),
                    "Event Type": "Early Close"
                }
            )

    # Building tables that will become Excel worksheets.
    summary_df = pd.DataFrame(
        summary_records,
        columns=[
            "Exchange",
            "Start Date",
            "End Date",
            "Full-day Holidays",
            "Early Closes"
        ]
    )

    events_df = pd.DataFrame(
        event_records,
        columns=[
            "Exchange",
            "Date",
            "Event Type"
        ]
    )

    if not events_df.empty:
        events_df = events_df.sort_values(
            by=["Exchange", "Date", "Event Type"]
        ).reset_index(drop=True)

    # Keeping a record of calendars that could not be processed.
    failed_df = pd.DataFrame(
        failed_exchanges,
        columns=["Exchange", "Error"]
    )

    # Checking the requested filename and making sure its folder exists.
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
        datetime_format="yyyy-mm-dd",
    ) as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        events_df.to_excel(writer, sheet_name="Calendar Events", index=False)
        failed_df.to_excel(writer, sheet_name="Failed Calendars", index=False)

        # Apply the same header, filtering, and layout to all three sheets.
        for worksheet in writer.sheets.values():
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.sheet_view.showGridLines = False

            for cell in worksheet[1]:
                cell.fill = PatternFill(fill_type="solid", fgColor="17365D")
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center")

            for column in worksheet.columns:
                worksheet.column_dimensions[column[0].column_letter].width = 23

        # Display dates consistently when opening the Excel workbooks.
        for row in writer.sheets["Summary"].iter_rows(min_row=2):
            row[1].number_format = "yyyy-mm-dd"
            row[2].number_format = "yyyy-mm-dd"

        for row in writer.sheets["Calendar Events"].iter_rows(min_row=2):
            row[1].number_format = "yyyy-mm-dd"

        # Leave enough room to read calendar error messages.
        writer.sheets["Failed Calendars"].column_dimensions["B"].width = 65

    print("\n" + "=" * 60)
    print("MARKET HOLIDAY CALENDAR")
    print("=" * 60)
    print(f"Exchange calendars: {len(summary_df)}")
    print(f"Full-day holidays:  {summary_df['Full-day Holidays'].sum()}")
    print(f"Early closes:       {summary_df['Early Closes'].sum()}")
    print(f"Failed calendars:   {len(failed_df)}")
    print(f"\nCompleted Excel workbook saved to: {output_path.resolve()}")

    if failed_exchanges:
        print("\nCalendar generation completed with failures; see 'Failed Calendars'.")
    else:
        print("\nMarket holiday calendar completed successfully.")


if __name__ == "__main__":
    main()
